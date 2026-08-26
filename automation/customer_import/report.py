from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


TABLE_LABELS = {
    "orders": "销售接单明细（订单头）", "activePlanItems": "月度计划（旧库）", "salesOrderLines": "数据中心-订单表",
    "inboundRows": "数据中心-入库表", "outboundRows": "数据中心-出库表", "salesOrders": "Planning 销售订单",
    "planItems": "月度计划（KDOS）", "orderSchedules": "营销中心-订单排期", "weeklyPlanItems": "主计划-周计划",
    "processProgressRows": "主计划-工序关联", "workReports": "主计划-报工表",
}


def create_report(output_dir: Path, customer_label: str, snapshots: list[dict[str, Any]], results: list[dict[str, Any]], dry_run: bool) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = output_dir / f"{customer_label}_客户数据导入清单_{timestamp}.xlsx"
    workbook = Workbook(); sheet = workbook.active; sheet.title = "导入表清单"
    sheet.append(["数据库", "数据表/报表", "写入行数", "处理状态", "说明"])
    if dry_run:
        sheet.append(["—", "—", 0, "仅预检", "未向 KDOS 写入数据"])
    for result in results:
        for database_key in ("legacy", "kdos"):
            database_label = "four_department_tracker" if database_key == "legacy" else "kdos"
            values = result.get(database_key, {})
            for key, value in values.items():
                if key in TABLE_LABELS:
                    note = "由源数据直接写入或聚合生成"
                    if key in {"weeklyPlanItems", "workReports"} and int(value or 0) == 0:
                        note = "按业务规则保留为手工同步，不伪造数据"
                    sheet.append([database_label, TABLE_LABELS[key], int(value or 0), "成功", note])
    source_sheet = workbook.create_sheet("源数据统计")
    source_sheet.append(["来源", "账套", "范围", "订单原始明细", "入库流水", "出库流水", "抽取时间", "幂等键"])
    for snapshot in snapshots:
        movements = snapshot.get("movements", [])
        source_sheet.append([
            snapshot["source"].upper(), snapshot["sourceDatabase"], snapshot["scope"].get("customerCode", "全部客户"), len(snapshot["orders"]),
            sum(1 for row in movements if row["direction"] == "INBOUND"), sum(1 for row in movements if row["direction"] == "OUTBOUND"),
            snapshot["extractedAt"], snapshot["idempotencyKey"],
        ])
    notes = workbook.create_sheet("口径与例外")
    notes.append(["项目", "说明"])
    notes.append(["T+ 出入库方向", "已按实际账套核验：方向值1=入库，方向值0=出库。"])
    notes.append(["计划范围", "订单表保留全部源明细；月度计划和订单排期只纳入已审核、未取消、未关闭且数量大于0的明细。"])
    notes.append(["客户要求交期", "T+ 当前没有可靠字段，未伪造；T+ 明细 deliveryDate 写入产前评审交期。"])
    notes.append(["周计划/报工表", "客户数据导入本身不伪造周排期或报工；需要时由业务人员通过系统的手工同步命令生成。"])
    notes.append(["业务人员对应表", "保留现有通讯录映射；A027 已存在对应关系，本次没有覆盖人工维护结果。"])
    for current in workbook.worksheets:
        current.freeze_panes = "A2"; current.auto_filter.ref = current.dimensions
        for cell in current[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4B70")
            cell.alignment = Alignment(horizontal="center")
        for column in current.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
            current.column_dimensions[column[0].column_letter].width = max(width, 12)
    workbook.save(path)
    return path
