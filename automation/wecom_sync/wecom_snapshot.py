import hashlib
import json
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from basic_code import WeChatPusher, export_contacts
from openpyxl import load_workbook


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _json_list(value: object) -> list[str]:
    text = _text(value)
    if not text:
        return []
    try:
        parsed = json.loads(text)
        return [str(item) for item in parsed] if isinstance(parsed, list) else []
    except json.JSONDecodeError:
        return [part.strip() for part in text.split("|") if part.strip()]


def _enabled(row: dict[str, object]) -> bool:
    enable = _text(row.get("enable")).lower()
    status = _text(row.get("status")).lower()
    return enable not in {"0", "false", "否", "停用", "禁用"} and status not in {"2", "4", "5"}


def _leader_in_department(value: object) -> bool:
    return _text(value).lower() in {"1", "true", "yes", "y", "是", "负责人"}


def _ignored(path: list[str]) -> bool:
    return (len(path) >= 1 and path[0] == "其他") or (len(path) >= 2 and path[1] == "其他")


def export_live_snapshot(output: Path) -> tuple[Path, dict[str, str]]:
    pusher = WeChatPusher()
    exported = export_contacts(output, pusher=pusher)
    return exported, pusher.get_departments()


def parse_snapshot(workbook_path: Path, department_paths: dict[str, str] | None = None) -> dict:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    headers = [_text(value) for value in next(rows)]
    grouped: OrderedDict[str, dict] = OrderedDict()
    discovered_paths: OrderedDict[str, list[str]] = OrderedDict()
    for values in rows:
        row = dict(zip(headers, values))
        user_id = _text(row.get("userid"))
        if not user_id:
            continue
        path: list[str] = []
        for index in range(1, 20):
            part = _text(row.get(f"部门{index}"))
            if not part:
                continue
            if not path or path[-1] != part:
                path.append(part)
        if not path or _ignored(path):
            continue
        department_id = _text(row.get("department_id"))
        if department_id:
            discovered_paths.setdefault(department_id, path)
        entry = grouped.setdefault(user_id, {
            "wechatUserId": user_id,
            "employeeNo": _text(row.get("工号")) or None,
            "name": _text(row.get("name")) or user_id,
            "position": _text(row.get("position")) or None,
            "telephone": _text(row.get("telephone")) or _text(row.get("mobile")) or None,
            "mobile": _text(row.get("mobile")) or _text(row.get("telephone")) or None,
            "email": _text(row.get("email")) or None,
            "alias": _text(row.get("alias")) or None,
            "gender": _text(row.get("gender")) or None,
            "directLeaders": [],
            "departmentLeaderExternalIds": [],
            "departmentPaths": [],
            "enabled": _enabled(row),
        })
        if path not in entry["departmentPaths"]:
            entry["departmentPaths"].append(path)
        entry["directLeaders"] = sorted(set(entry["directLeaders"] + _json_list(row.get("direct_leader"))))
        if department_id and _leader_in_department(row.get("is_leader_in_dept")) and department_id not in entry["departmentLeaderExternalIds"]:
            entry["departmentLeaderExternalIds"].append(department_id)

    raw_paths = department_paths or {department_id: " / ".join(path) for department_id, path in discovered_paths.items()}
    normalized_paths: OrderedDict[str, list[str]] = OrderedDict()
    for department_id, raw_path in raw_paths.items():
        path = [part.strip() for part in raw_path.split(" / ") if part.strip()]
        if path and not _ignored(path):
            normalized_paths[str(department_id)] = path
    path_ids = {" / ".join(path): department_id for department_id, path in normalized_paths.items()}
    departments = []
    for sort_order, (department_id, path) in enumerate(normalized_paths.items()):
        parent_id = path_ids.get(" / ".join(path[:-1])) if len(path) > 1 else None
        departments.append({
            "externalId": department_id,
            "name": path[-1],
            "parentExternalId": parent_id,
            "path": path,
            "sortOrder": sort_order,
        })
    digest = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    return {
        "capturedAt": datetime.now().astimezone().isoformat(),
        "sourceHash": digest,
        "departments": departments,
        "contacts": list(grouped.values()),
    }
